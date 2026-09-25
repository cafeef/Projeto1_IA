"""Gera docs/RELATORIO.docx (modelo UTFPR) a partir de docs/RELATORIO.md.

Uso, a partir da raiz do repositório:
    python scripts/gerar_relatorio_docx.py            # .docx + prévia em PDF
    python scripts/gerar_relatorio_docx.py --no-pdf   # só o .docx

O modelo oficial (docs/modelo/modelo_utfpr_sem_licenca.docx) fornece estilos,
margens, capa, folha de rosto e cabeçalho com paginação. Os elementos
pré-textuais ficam em PRE_TEXTUAL abaixo; o corpo vem do Markdown:

- "## N. Título" / "### N.N Título" viram seções numeradas pelo próprio modelo;
- "Quadro N – ...", "Tabela N – ...", "Gráfico N – ...", "Figura N – ..." viram
  legendas; "Fonte: ..." vira a fonte da ilustração;
- tabelas Markdown, imagens ![](...), listas, ```formula (equação numerada) e
  ```text (bloco monoespaçado) são convertidos; citações ">" são ignoradas.

Com LibreOffice (soffice) e PyMuPDF disponíveis, o script renderiza o documento,
descobre em que página cada seção e ilustração caiu e preenche o sumário, as
listas e o número da primeira página do texto. O Word ainda atualiza os campos
ao abrir o arquivo (updateFields), então os números se ajustam à paginação dele.
"""

import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "docs" / "modelo" / "modelo_utfpr_sem_licenca.docx"
SOURCE = ROOT / "docs" / "RELATORIO.md"
OUTPUT = ROOT / "docs" / "RELATORIO.docx"

# ------------------------------------------------------------ pré-textuais

PRE_TEXTUAL = {
    # Nomes completos, em ordem alfabética.
    "autores": [
        "ANDRÉ MARTINS DA SILVA",
        "BRYAN WILLIAN DE OLIVEIRA BETIM",
        "FERNANDA PACHECO BENTO",
        "OTÁVIO LUIS HRUBA",
    ],
    "titulo": "AGENTE DE COMBATE À DENGUE: RESOLUÇÃO DE PROBLEMA POR MEIO DE ALGORITMOS DE BUSCA",
    "cidade": "PONTA GROSSA",
    "ano": "2026",
    "natureza": (
        "Relatório técnico apresentado como requisito parcial para aprovação na disciplina "
        "de Inteligência Artificial, do Curso de Bacharelado em Ciência da Computação da "
        "Universidade Tecnológica Federal do Paraná, Campus Ponta Grossa."
    ),
    "professor": "Professora: Helyane Bronoski Borges.",
    "resumo": (
        "Este trabalho apresenta um ambiente de simulação em que um usuário humano e um "
        "agente inteligente resolvem a mesma missão: partir de uma posição inicial e alcançar "
        "um foco de dengue em uma grade bidimensional com obstáculos e terrenos de custos "
        "diferentes. O ambiente foi modelado como um problema de busca em grafo implícito, e "
        "foram implementados, de forma autoral, a Busca em Largura, a Busca em Profundidade, a "
        "Busca Gulosa e o A*. Os dois últimos usam como heurística a distância Manhattan "
        "multiplicada pelo menor custo de movimento, que é admissível e consistente. A "
        "interface, desenvolvida em Python com a biblioteca Arcade, exibe os estados "
        "explorados, os caminhos e as métricas de cada participante, e apresenta ao final uma "
        "mensagem educativa sobre o criadouro encontrado. Foram realizadas 15 execuções em "
        "três cenários de complexidade crescente, sendo três humanas e doze algorítmicas. A "
        "Busca em Largura encontrou sempre o menor número de passos, mas, no cenário complexo, "
        "seu caminho custou 65, contra 21 do A*, que obteve o menor custo em todos os "
        "cenários expandindo poucos estados a mais que a Busca Gulosa. A Busca em Profundidade "
        "produziu os piores caminhos. O usuário alcançou o custo ótimo nos três cenários "
        "oficiais, mas foi superado pelo agente em um labirinto com desvios e no cenário sem "
        "solução. Conclui-se que o A* é a estratégia mais adequada ao problema."
    ),
    "palavras_chave": "inteligência artificial; algoritmos de busca; heurística; A*; dengue.",
    "abstract": (
        "This work presents a simulation environment in which a human user and an intelligent "
        "agent solve the same mission: starting from an initial position and reaching a dengue "
        "breeding site on a two-dimensional grid with obstacles and terrains of different "
        "costs. The environment was modeled as a search problem over an implicit graph, and "
        "Breadth-First Search, Depth-First Search, Greedy Search and A* were implemented from "
        "scratch. The last two use as heuristic the Manhattan distance multiplied by the "
        "lowest movement cost, which is admissible and consistent. The interface, developed in "
        "Python with the Arcade library, displays the explored states, the paths and the "
        "metrics of each participant, and shows an educational message about the breeding "
        "site found. Fifteen runs were carried out in three scenarios of increasing "
        "complexity, three by the user and twelve by the algorithms. Breadth-First Search "
        "always found the fewest steps, but in the complex scenario its path cost 65, against "
        "21 for A*, which obtained the lowest cost in every scenario while expanding only a few "
        "more states than Greedy Search. Depth-First Search produced the worst paths. The user "
        "reached the optimal cost in the three official scenarios, but was outperformed by the "
        "agent in a maze with detours and in the scenario without a solution. A* is therefore "
        "the most suitable strategy for the problem."
    ),
    "keywords": "artificial intelligence; search algorithms; heuristics; A*; dengue.",
    "siglas": [
        ("BFS", "Busca em Largura (do inglês Breadth-First Search)"),
        ("CSV", "Valores separados por vírgula (do inglês Comma-Separated Values)"),
        ("DFS", "Busca em Profundidade (do inglês Depth-First Search)"),
        ("FIFO", "Primeiro a entrar, primeiro a sair (do inglês First In, First Out)"),
        ("JSON", "Notação de objetos JavaScript (do inglês JavaScript Object Notation)"),
        ("LIFO", "Último a entrar, primeiro a sair (do inglês Last In, First Out)"),
        ("OMS", "Organização Mundial da Saúde"),
        ("UTFPR", "Universidade Tecnológica Federal do Paraná"),
    ],
    "simbolos": [
        ("n", "Estado (célula) da grade"),
        ("g(n)", "Custo do caminho da origem até n"),
        ("h(n)", "Custo estimado de n até o foco (heurística)"),
        ("h*(n)", "Custo real do menor caminho de n até o foco"),
        ("f(n)", "Função de avaliação que ordena a fronteira"),
        ("c(n, n')", "Custo do passo de n para o vizinho n'"),
        ("c_min", "Menor custo de movimento entre os terrenos transitáveis"),
    ],
}

# Estilos do modelo (os ids são numéricos neste arquivo).
S_TEXT = "936"  # Texto do Trabalho
S_H1, S_H2, S_H3 = "942", "895", "896"
S_PRE_TITLE = "934"  # Titulo 6 (RESUMO, SUMÁRIO...)
S_ABSTRACT = "937"  # Formatação do resumo
S_CAPTION = "958"
S_FIGURE = "959"  # Parágrafo para Ilustrações
S_SOURCE = "960"  # Fonte das Ilustrações
S_CELL_HEAD, S_CELL_BODY = "954", "955"
S_REF_TITLE, S_REF = "978", "947"
S_TOC = {1: "973", 2: "974", 3: "975", 6: "979"}
S_LIST_OF = "972"  # table of figures
S_NORMAL = "893"
S_HYPERLINK = "919"
BULLET_NUM_ID = "49"
TEXT_WIDTH = 9071  # 16 cm em twips (A4 com margens 3/2 cm)

CAPTION_RE = re.compile(r"^(Quadro|Tabela|Gráfico|Figura) (\d+) – (.+)$")
FIGURE_KINDS = ("Figura", "Gráfico", "Quadro")  # vão para a lista de ilustrações


# ----------------------------------------------------------- XML helpers


def el(tag: str, **attrs) -> OxmlElement:
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(f"w:{key}"), str(value))
    return node


def sub(parent, tag: str, **attrs):
    node = el(tag, **attrs)
    parent.append(node)
    return node


def paragraph(style: str, **ppr) -> OxmlElement:
    p = el("w:p")
    pPr = sub(p, "w:pPr")
    sub(pPr, "w:pStyle", val=style)
    if "keep_next" in ppr:
        sub(pPr, "w:keepNext", val=1 if ppr["keep_next"] else 0)
    if "num" in ppr:
        numPr = sub(pPr, "w:numPr")
        sub(numPr, "w:ilvl", val=0)
        sub(numPr, "w:numId", val=ppr["num"])
    if "tabs" in ppr:
        tabs = sub(pPr, "w:tabs")
        for kind, pos, leader in ppr["tabs"]:
            sub(tabs, "w:tab", val=kind, pos=pos, leader=leader)
    if "spacing" in ppr:
        sub(pPr, "w:spacing", **ppr["spacing"])
    if "ind" in ppr:
        sub(pPr, "w:ind", **ppr["ind"])
    if "jc" in ppr:
        sub(pPr, "w:jc", val=ppr["jc"])
    return p


def run(text: str = "", bold=False, italic=False, mono=False, size=None, style=None):
    r = el("w:r")
    rPr = sub(r, "w:rPr")
    if style:
        sub(rPr, "w:rStyle", val=style)
    if mono:
        sub(rPr, "w:rFonts", ascii="Courier New", hAnsi="Courier New", cs="Courier New")
    if bold:
        sub(rPr, "w:b")
    if italic:
        sub(rPr, "w:i")
    if size:
        sub(rPr, "w:sz", val=size)
        sub(rPr, "w:szCs", val=size)
    if text:
        t = sub(r, "w:t")
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return r


def tab_run():
    r = el("w:r")
    sub(r, "w:tab")
    return r


def page_break_run():
    r = el("w:r")
    sub(r, "w:br", type="page")
    return r


def field_runs(instr: str, cached: str = "", parts=("begin", "instr", "separate", "text", "end")):
    runs = []
    for part in parts:
        if part == "instr":
            r = el("w:r")
            t = sub(r, "w:instrText")
            t.text = f" {instr} "
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        elif part == "text":
            r = run(cached)
        else:
            r = el("w:r")
            sub(r, "w:fldChar", fldCharType=part)
        runs.append(r)
    return runs


class Bookmarks:
    def __init__(self) -> None:
        self.next_id = 5000

    def wrap(self, p, name: str) -> None:
        """Marca todo o conteúdo do parágrafo com um bookmark."""
        self.next_id += 1
        start = el("w:bookmarkStart", id=self.next_id, name=name)
        end = el("w:bookmarkEnd", id=self.next_id)
        pPr = p.find(qn("w:pPr"))
        p.insert(list(p).index(pPr) + 1, start)
        p.append(end)


# ------------------------------------------------------ texto com markdown

INLINE_RE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")
STAR = ""  # guarda "\*" durante o parse


def inline_runs(text: str, bold=False, italic=False, size=None, mono_size=None):
    text = text.replace("\\*", STAR)
    runs = []
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append(run(_star(text[pos:m.start()]), bold, italic, size=size))
        token = m.group(0)
        if token.startswith("`"):
            runs.append(run(_star(token[1:-1]), bold, italic, mono=True, size=mono_size or size))
        elif token.startswith("**"):
            runs += inline_runs(token[2:-2].replace(STAR, "\\*"), True, italic, size, mono_size)
        else:
            runs += inline_runs(token[1:-1].replace(STAR, "\\*"), bold, True, size, mono_size)
        pos = m.end()
    if pos < len(text):
        runs.append(run(_star(text[pos:]), bold, italic, size=size))
    return runs


def _star(text: str) -> str:
    return text.replace(STAR, "*")


def plain(text: str) -> str:
    """Texto sem marcação (para sumário e legendas)."""
    text = text.replace("\\*", STAR)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])", r"\1", text)
    return _star(text)


# ------------------------------------------------------------ Markdown


@dataclass
class Block:
    kind: str  # h1 h2 h3 para bullets alineas table code formula image caption source
    text: str = ""
    items: list = field(default_factory=list)
    lang: str = ""


def parse_markdown(md: str) -> list[Block]:
    lines = md.splitlines()
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.startswith("# ") or line.startswith(">"):
            i += 1
            continue
        if line.startswith("```"):
            lang = line[3:].strip()
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            blocks.append(Block("formula" if lang == "formula" else "code", items=lines[i + 1:j]))
            i = j + 1
            continue
        if line.startswith("## ") or line.startswith("### "):
            level = 2 if line.startswith("## ") else 3
            title = re.sub(r"^[\d.]+\s+", "", line[level + 1:].strip())
            blocks.append(Block("h1" if level == 2 else "h2", plain(title)))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            aligns = ["right" if c.endswith(":") else "left" for c in rows[1]]
            blocks.append(Block("table", items=[rows[0], aligns, rows[2:]]))
            continue
        if re.match(r"^!\[.*\]\((.+)\)$", line.strip()):
            path = re.match(r"^!\[.*\]\((.+)\)$", line.strip()).group(1)
            blocks.append(Block("image", path))
            i += 1
            continue
        list_match = re.match(r"^(- |\d+\. )", line)
        if list_match:
            ordered = line[0].isdigit()
            items = []
            while i < len(lines) and lines[i].strip():
                if re.match(r"^(- |\d+\. )", lines[i]):
                    items.append(re.sub(r"^(- |\d+\. )", "", lines[i]).strip())
                else:
                    items[-1] += " " + lines[i].strip()
                i += 1
            blocks.append(Block("alineas" if ordered else "bullets", items=items))
            continue
        # parágrafo comum: junta as linhas até a linha em branco
        chunk = []
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("|", "```", "#", "- ")):
            chunk.append(lines[i].strip())
            i += 1
        text = " ".join(chunk)
        caption = CAPTION_RE.match(text)
        if caption:
            blocks.append(Block("caption", text, items=list(caption.groups())))
        elif text.startswith("Fonte:"):
            blocks.append(Block("source", text))
        elif re.match(r"^\*\*[^*]+\*\*$", chunk[0]) and len(chunk) > 1:
            # "**Pergunta?**" seguida da resposta: dois parágrafos
            blocks.append(Block("question", chunk[0][2:-2]))
            blocks.append(Block("para", " ".join(chunk[1:])))
        else:
            blocks.append(Block("para", text))
    return blocks


# ------------------------------------------------------------- tabelas


def table_xml(header, aligns, rows, closed: bool):
    """Quadro: todas as bordas. Tabela: só bordas horizontais (IBGE/ABNT)."""
    weights = []
    for col in range(len(header)):
        longest = max(len(plain(r[col])) for r in rows)
        header_word = max(len(w) for w in plain(header[col]).split())
        weights.append(max(min(longest, 30), header_word) + 3)
    total = sum(weights)
    widths = [int(TEXT_WIDTH * w / total) for w in weights]
    widths[-1] += TEXT_WIDTH - sum(widths)

    tbl = el("w:tbl")
    tblPr = sub(tbl, "w:tblPr")
    sub(tblPr, "w:jc", val="center")
    sub(tblPr, "w:tblW", w=TEXT_WIDTH, type="dxa")
    borders = sub(tblPr, "w:tblBorders")
    thick = {"val": "single", "sz": 12, "space": 0, "color": "000000"}
    thin = {"val": "single", "sz": 4, "space": 0, "color": "000000"}
    none = {"val": "nil"}
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if closed:
            attrs = thin
        elif side in ("top", "bottom"):
            attrs = thick
        else:
            attrs = none
        sub(borders, f"w:{side}", **attrs)
    sub(tblPr, "w:tblLayout", type="fixed")
    margins = sub(tblPr, "w:tblCellMar")
    sub(margins, "w:left", w=60, type="dxa")
    sub(margins, "w:right", w=60, type="dxa")
    grid = sub(tbl, "w:tblGrid")
    for w in widths:
        sub(grid, "w:gridCol", w=w)

    for r_index, cells in enumerate([header] + rows):
        tr = sub(tbl, "w:tr")
        trPr = sub(tr, "w:trPr")
        sub(trPr, "w:cantSplit")
        if r_index == 0:
            sub(trPr, "w:tblHeader")
        for c_index, cell in enumerate(cells):
            tc = sub(tr, "w:tc")
            tcPr = sub(tc, "w:tcPr")
            sub(tcPr, "w:tcW", w=widths[c_index], type="dxa")
            if r_index == 0 and not closed:
                b = sub(tcPr, "w:tcBorders")
                sub(b, "w:bottom", **thin)
            sub(tcPr, "w:vAlign", val="center")
            if r_index == 0:
                p = paragraph(S_CELL_HEAD)
            else:
                p = paragraph(S_CELL_BODY, jc="center" if aligns[c_index] == "right" else "left")
            for r in inline_runs(cell, bold=r_index == 0, size=20, mono_size=18):
                p.append(r)
            tc.append(p)
    return tbl


# ------------------------------------------------------------ montagem


@dataclass
class Entry:
    """Item de sumário ou lista: preenchido com a página depois da renderização."""

    level: int  # 1, 2, 3 (seções), 6 (REFERÊNCIAS) ou 0 (legenda)
    number: str
    title: str
    bookmark: str
    kind: str = ""  # Figura, Gráfico, Quadro ou Tabela, para legendas
    page: int = 0


class Builder:
    def __init__(self, doc: Document) -> None:
        self.doc = doc
        self.out: list = []
        self.entries: list[Entry] = []
        self.bookmarks = Bookmarks()
        self.counters = [0, 0, 0]
        self.caption_kind = ""
        self.first_h1 = True

    def _bm(self) -> str:
        return f"_Toc9{len(self.entries):07d}"

    def heading(self, level: int, title: str) -> None:
        style = {1: S_H1, 2: S_H2, 3: S_H3}[level]
        self.counters[level - 1] += 1
        for k in range(level, 3):
            self.counters[k] = 0
        number = ".".join(str(c) for c in self.counters[:level])
        p = paragraph(style)
        if level == 1 and not self.first_h1:
            p.append(page_break_run())
        self.first_h1 = False
        p.append(run(title))
        name = self._bm()
        self.bookmarks.wrap(p, name)
        self.entries.append(Entry(level, number, title, name))
        self.out.append(p)

    def references_title(self) -> None:
        p = paragraph(S_REF_TITLE)
        p.append(page_break_run())
        p.append(run("REFERÊNCIAS"))
        name = self._bm()
        self.bookmarks.wrap(p, name)
        self.entries.append(Entry(6, "", "REFERÊNCIAS", name))
        self.out.append(p)

    def text(self, text: str, style=S_TEXT, **ppr) -> None:
        p = paragraph(style, **ppr)
        for r in inline_runs(text, mono_size=22):
            p.append(r)
        self.out.append(p)

    def question(self, text: str) -> None:
        p = paragraph(S_TEXT, keep_next=True)
        for r in inline_runs(text, bold=True, mono_size=22):
            p.append(r)
        self.out.append(p)

    def bullets(self, items) -> None:
        for item in items:
            self.text(item, num=BULLET_NUM_ID, ind={"left": 1134, "hanging": 283})

    def alineas(self, items) -> None:
        for n, item in enumerate(items):
            self.text(f"{chr(ord('a') + n)}) {item}", ind={"left": 1134, "hanging": 425})

    def caption(self, kind: str, number: str, title: str) -> None:
        self.caption_kind = kind
        p = paragraph(S_CAPTION, keep_next=True)
        p.append(run(f"{kind} "))
        for r in field_runs(f"SEQ {kind} \\* ARABIC", number):
            p.append(r)
        p.append(run(f" - {plain(title)}"))
        name = self._bm()
        self.bookmarks.wrap(p, name)
        self.entries.append(Entry(0, number, plain(title), name, kind))
        self.out.append(p)

    def source(self, text: str) -> None:
        # O estilo do modelo encadeia "manter com o próximo" até o texto seguinte,
        # o que arrasta vários gráficos juntos; a fonte encerra a ilustração.
        p = paragraph(S_SOURCE, keep_next=False)
        for r in inline_runs(text, size=20):
            p.append(r)
        self.out.append(p)

    def table(self, header, aligns, rows) -> None:
        self.out.append(table_xml(header, aligns, rows, closed=self.caption_kind != "Tabela"))

    def image(self, rel_path: str) -> None:
        path = (SOURCE.parent / rel_path).resolve()
        para = self.doc.add_paragraph()
        pPr = para._p.get_or_add_pPr()
        pPr.insert(0, el("w:pStyle", val=S_FIGURE))
        para.add_run().add_picture(str(path), width=Cm(15.5))
        self.out.append(para._p)

    def code(self, lines) -> None:
        for n, line in enumerate(lines):
            p = paragraph(S_FIGURE, jc="left", keep_next=n < len(lines) - 1)
            p.append(run(line, mono=True, size=14))
            self.out.append(p)

    def formula(self, lines, number: int) -> None:
        p = paragraph(
            S_NORMAL,
            tabs=[("center", TEXT_WIDTH // 2, "none"), ("right", TEXT_WIDTH, "none")],
            spacing={"before": 240, "after": 240},
            ind={"left": 0, "firstLine": 0},
        )
        p.append(tab_run())
        p.append(run(" ".join(lines)))
        p.append(tab_run())
        p.append(run(f"({number})"))
        self.out.append(p)

    def reference(self, text: str) -> None:
        p = paragraph(S_REF)
        for r in inline_runs(text):
            p.append(r)
        self.out.append(p)

    def build_body(self, blocks: list[Block]) -> None:
        in_refs = False
        equations = 0
        for block in blocks:
            if block.kind == "h1":
                if block.text.lower().startswith("referências"):
                    in_refs = True
                    self.references_title()
                else:
                    self.heading(1, block.text)
            elif block.kind == "h2":
                self.heading(2, block.text)
            elif in_refs and block.kind == "para":
                self.reference(block.text)
            elif block.kind == "para":
                self.text(block.text)
            elif block.kind == "question":
                self.question(block.text)
            elif block.kind == "bullets":
                self.bullets(block.items)
            elif block.kind == "alineas":
                self.alineas(block.items)
            elif block.kind == "caption":
                self.caption(*block.items)
            elif block.kind == "source":
                self.source(block.text)
            elif block.kind == "table":
                self.table(*block.items)
            elif block.kind == "image":
                self.image(block.text)
            elif block.kind == "code":
                self.code(block.items)
            elif block.kind == "formula":
                equations += 1
                self.formula(block.items, equations)


# ------------------------------------------------------- pré-textuais


def set_text(p, lines: list[str]) -> None:
    """Troca o texto de um parágrafo do modelo mantendo formatação do 1º run."""
    rPr = None
    for r in p.findall(qn("w:r")):
        if r.find(qn("w:t")) is not None and r.find(qn("w:rPr")) is not None:
            rPr = r.find(qn("w:rPr"))
            break
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    for n, line in enumerate(lines):
        r = el("w:r")
        if rPr is not None:
            props = copy.deepcopy(rPr)
            for color in props.findall(qn("w:color")):
                props.remove(color)
            r.append(props)
        if n:
            sub(r, "w:br")
        t = sub(r, "w:t")
        t.text = line
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        p.append(r)


def pre_title(text: str):
    p = paragraph(S_PRE_TITLE)
    p.append(page_break_run())
    p.append(run(text))
    return p


def abstract_page(title: str, body: str, label: str, words: str) -> list:
    p = paragraph(S_ABSTRACT)
    p.append(run(body))
    k = paragraph(S_ABSTRACT)
    k.append(run(f"{label}: ", bold=True))
    k.append(run(words))
    return [pre_title(title), p, k]


def two_column_list(template_tbl, rows):
    """Reaproveita a tabela sem bordas do modelo (siglas/símbolos)."""
    tbl = copy.deepcopy(template_tbl)
    trs = tbl.findall(qn("w:tr"))
    proto = trs[0]
    for tr in trs:
        tbl.remove(tr)
    for key, value in rows:
        tr = copy.deepcopy(proto)
        cells = tr.findall(qn("w:tc"))
        set_text(cells[0].find(qn("w:p")), [key])
        set_text(cells[1].find(qn("w:p")), [value])
        tbl.append(tr)
    return tbl


def toc_entry(style: str, parts: list, bookmark: str, page: int, tabs=None):
    p = paragraph(style, tabs=tabs) if tabs else paragraph(style)
    link = el("w:hyperlink")
    link.set(qn("w:anchor"), bookmark)
    link.set(qn("w:history"), "1")
    for part in parts:
        link.append(tab_run() if part == "\t" else run(part))
    link.append(tab_run())
    for r in field_runs(f"PAGEREF {bookmark} \\h", str(page)):
        link.append(r)
    p.append(link)
    return p


def field_list(instr: str, entry_paragraphs: list) -> list:
    """Envolve os itens num campo TOC (Word regenera ao atualizar)."""
    if not entry_paragraphs:
        return []
    first, last = entry_paragraphs[0], entry_paragraphs[-1]
    pPr = first.find(qn("w:pPr"))
    for n, r in enumerate(field_runs(instr, parts=("begin", "instr", "separate"))):
        first.insert(list(first).index(pPr) + 1 + n, r)
    last.append(field_runs("", parts=("end",))[0])
    return entry_paragraphs


def lists_and_summary(entries: list[Entry]) -> list:
    out = []
    list_tabs = [("right", TEXT_WIDTH, "dot")]

    out.append(pre_title("LISTA DE ILUSTRAÇÕES"))
    for kind in FIGURE_KINDS:
        items = [
            toc_entry(S_LIST_OF, [f"{kind} {e.number} - {e.title}"], e.bookmark, e.page, list_tabs)
            for e in entries if e.kind == kind
        ]
        out += field_list(f'TOC \\h \\z \\c "{kind}"', items)

    out.append(pre_title("LISTA DE TABELAS"))
    items = [
        toc_entry(S_LIST_OF, [f"Tabela {e.number} - {e.title}"], e.bookmark, e.page, list_tabs)
        for e in entries if e.kind == "Tabela"
    ]
    out += field_list('TOC \\h \\z \\c "Tabela"', items)
    return out


def summary(entries: list[Entry]) -> list:
    out = [pre_title("SUMÁRIO")]
    items = []
    for e in entries:
        if e.level == 6:
            items.append(toc_entry(S_TOC[6], [e.title], e.bookmark, e.page))
        elif e.level in (1, 2, 3):
            items.append(toc_entry(S_TOC[e.level], [e.number, "\t", e.title], e.bookmark, e.page))
    out += field_list(
        'TOC \\o "2-5" \\h \\z \\t "Titulo 1;1;Título REFERÊNCIAS;6"', items
    )
    return out


def build(entries_pages: dict | None, first_page: int) -> tuple[Document, list[Entry]]:
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    kids = list(body)
    texts = ["".join(t.text or "" for t in k.iter(qn("w:t"))) for k in kids]
    assert texts[0].startswith("UNIVERSIDADE TECNOLÓGICA"), "modelo inesperado"
    assert texts[195] == "Resumo" and texts[267] == "" and texts[227].startswith("ABNT")

    info = PRE_TEXTUAL
    # capa
    set_text(kids[6], info["autores"])
    set_text(kids[14], [info["titulo"]])
    set_text(kids[30], [info["cidade"]])
    set_text(kids[31], [info["ano"]])
    # folha de rosto
    set_text(kids[32], info["autores"])
    set_text(kids[42], [info["titulo"]])
    set_text(kids[49], [info["natureza"]])
    set_text(kids[50], [info["professor"]])
    set_text(kids[63], [info["cidade"]])
    set_text(kids[64], [info["ano"]])
    # Capa: as 3 linhas extras dos autores saem das linhas em branco 7-9.
    # Folha de rosto: sai a nota do modelo, o título traduzido e o coorientador.
    assert all(texts[n] == "" for n in (7, 8, 9))
    cover = [k for n, k in enumerate(kids[:65]) if n not in (7, 8, 9, 44, 46, 51)]

    siglas_tbl, simbolos_tbl = kids[227], kids[231]
    section_break = kids[267]
    section_break.find(qn("w:pPr")).find(qn("w:pStyle")).set(qn("w:val"), S_NORMAL)
    final_sect = kids[-1]
    final_sect.find(qn("w:pgNumType")).set(qn("w:start"), str(first_page))

    for k in kids:
        body.remove(k)

    builder = Builder(doc)
    builder.build_body(parse_markdown(SOURCE.read_text(encoding="utf-8")))
    for e in builder.entries:
        e.page = (entries_pages or {}).get(e.bookmark, 0)

    pre = cover
    pre += abstract_page("RESUMO", info["resumo"], "Palavras-chave", info["palavras_chave"])
    pre += abstract_page("ABSTRACT", info["abstract"], "Keywords", info["keywords"])
    pre += lists_and_summary(builder.entries)
    pre += [pre_title("LISTA DE ABREVIATURAS E SIGLAS"), two_column_list(siglas_tbl, info["siglas"])]
    pre += [pre_title("LISTA DE SÍMBOLOS"), two_column_list(simbolos_tbl, info["simbolos"])]
    pre += summary(builder.entries)
    pre.append(section_break)

    for node in pre + builder.out:
        body.append(node)
    body.append(final_sect)

    # Imagens e gráfico de exemplo do modelo não são mais referenciados.
    xml = body.xml
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith(("/image", "/chart", "/hyperlink")) and f'"{rid}"' not in xml:
            doc.part.drop_rel(rid)

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = el("w:updateFields", val="true")
        anchor = settings.find(qn("w:footnotePr"))
        if anchor is not None:
            anchor.addprevious(update)
        else:
            settings.append(update)
    return doc, builder.entries


# ------------------------------------------------ paginação via LibreOffice


def render_pdf(docx: Path, out_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx)],
        check=True, capture_output=True, timeout=240,
    )
    return out_dir / (docx.stem + ".pdf")


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).casefold()
    return re.sub(r"\s+", " ", text)


def locate(pdf: Path, entries: list[Entry]) -> tuple[int, dict] | None:
    """Página física (0 = capa) de cada entrada e do início do texto."""
    try:
        import pymupdf
    except ImportError:
        return None
    pages = []
    body_start = None
    with pymupdf.open(pdf) as document:
        for index, page in enumerate(document):
            pages.append(norm(page.get_text()))
            top = [b[4].strip() for b in page.get_text("blocks") if b[1] < 100]
            if body_start is None and any(t.isdigit() for t in top):
                body_start = index
    if body_start is None:
        return None
    found = {}
    cursor = body_start
    for e in entries:
        if e.kind:
            pattern = re.escape(norm(f"{e.kind} {e.number} - {e.title}")[:40])
        elif e.level == 6:
            pattern = "referências"
        else:
            pattern = rf"(^| ){re.escape(e.number)} {re.escape(norm(e.title)[:30])}"
        for index in range(cursor, len(pages)):
            if re.search(pattern, pages[index]):
                found[e.bookmark] = index
                cursor = index
                break
    return body_start, found


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--no-pdf", action="store_true", help="não gera a prévia em PDF")
    args = parser.parse_args()

    doc, entries = build(None, 1)
    doc.save(OUTPUT)
    with tempfile.TemporaryDirectory() as tmp:
        pdf = render_pdf(OUTPUT, Path(tmp))
        located = locate(pdf, entries) if pdf else None
        if located is None:
            print("Sem LibreOffice/PyMuPDF: páginas ficam para o Word atualizar (F9).")
            print(f"Gerado: {OUTPUT.relative_to(ROOT)}")
            return
        body_start, pages = located
        # Capa não é contada; a folha de rosto é a página 1.
        doc, entries = build(pages, body_start)
        doc.save(OUTPUT)
        missing = [e.title for e in entries if e.bookmark not in pages]
        if missing:
            print("Aviso: páginas não localizadas para:", "; ".join(missing), file=sys.stderr)
        print(f"Gerado: {OUTPUT.relative_to(ROOT)} (texto começa na página {body_start})")
        if not args.no_pdf:
            pdf = render_pdf(OUTPUT, Path(tmp))
            shutil.copy(pdf, OUTPUT.with_suffix(".pdf"))
            print(f"Prévia: {OUTPUT.with_suffix('.pdf').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
